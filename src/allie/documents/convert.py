"""Turning an Excel workbook, a CSV, a Word file or a text file into the
text the document model reads (plan.md D35, spec of 2026-09-25 section 3.1).

A PDF is not converted: Gemini reads a PDF itself, page by page as an image,
tables and all - better than any text pulled out of it here. Nothing else a
folder holds is something Gemini reads, so it becomes Markdown: a table is
a table and a paragraph a paragraph.

**Plain on purpose** (the owner, 2026-09-25: findings 8 and 10 of the
review are not wanted). An Excel cell is its raw value - the 18 % of a VAT
rate is `0.18`, a date is `2026-08-14 00:00:00` - and a formula Excel never
calculated is empty; a Word file is its body, paragraphs and tables in
order, without its header and footer.

**Caps.** `MAX_ROWS` rows a table and `MAX_CHARS` characters a file, and
the cut is said in the text itself, so that the model can say it too.

The libraries are imported where they are used: `allie.config` reaches
this package for its defaults, and no command should pay for openpyxl or
lxml before a file is read.
"""

from __future__ import annotations

import csv
import io
from collections.abc import Iterable, Sequence
from pathlib import Path
from typing import Any

__all__ = [
    "CSV_DELIMITERS",
    "CUT",
    "MAX_CHARS",
    "MAX_ROWS",
    "ROWS_CUT",
    "SNIFF_CHARS",
    "capped",
    "cell_text",
    "csv_text",
    "decoded",
    "excel_text",
    "plain_text",
    "table",
    "word_text",
]

# A sheet of a few pages is a few dozen rows; two hundred leaves room and
# keeps an exported ledger from filling a question on its own.
MAX_ROWS = 200
# Ten pages of text are some 30 000 characters.
MAX_CHARS = 40_000
ROWS_CUT = "(This table has {total} rows; the first {limit} are above.)"
CUT = "(This file has {total} characters; the first {limit} are above.)"

# The separators a CSV is tried with: a Turkish Excel writes `;`, because
# `,` is its decimal point.
CSV_DELIMITERS = ";,\t|"
SNIFF_CHARS = 4_096


def cell_text(value: object) -> str:
    """A cell as text: its raw value, on one line; nothing for an empty one."""
    if value is None:
        return ""
    return " ".join(str(value).split())


def table(rows: Iterable[Sequence[str]]) -> str:
    """Rows as a Markdown table under the first row.

    Empty rows are left out and the columns end at the last one anything is
    in; a `|` in a cell is escaped so that it does not end the cell.
    """
    kept = [[_clean(cell) for cell in row] for row in rows]
    kept = [row for row in kept if any(row)]
    if not kept:
        return ""
    width = max(max(index + 1 for index, cell in enumerate(row) if cell) for row in kept)
    header, *body = kept
    lines = [_row(header, width), _row(["---"] * width, width)]
    lines.extend(_row(row, width) for row in body[:MAX_ROWS])
    if len(body) > MAX_ROWS:
        lines.append(ROWS_CUT.format(total=len(body), limit=MAX_ROWS))
    return "\n".join(lines)


def capped(text: str) -> str:
    """`text`, or its first `MAX_CHARS` characters and a line that says so."""
    if len(text) <= MAX_CHARS:
        return text
    return f"{text[:MAX_CHARS]}\n{CUT.format(total=len(text), limit=MAX_CHARS)}"


def decoded(raw: bytes) -> str:
    """Bytes as text: UTF-8, with or without its mark, else the Turkish
    Windows code page every older export uses."""
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("cp1254", errors="replace")


def excel_text(path: Path) -> str:
    """Every sheet with anything in it, under its name, as a table."""
    import openpyxl

    # `data_only`: the value Excel saved, not the formula behind it.
    book: Any = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        parts: list[str] = []
        for sheet in book.worksheets:
            rows = [
                [cell_text(value) for value in row] for row in sheet.iter_rows(values_only=True)
            ]
            text = table(rows)
            if text:
                parts.append(f"## {sheet.title}\n\n{text}")
        return capped("\n\n".join(parts))
    finally:
        book.close()


def csv_text(path: Path) -> str:
    """A CSV as one table, whichever separator and code page wrote it."""
    text = decoded(path.read_bytes())
    dialect: type[csv.Dialect]
    try:
        dialect = csv.Sniffer().sniff(text[:SNIFF_CHARS], delimiters=CSV_DELIMITERS)
    except csv.Error:
        # One column, or nothing the sniffer can be sure of: a comma.
        dialect = csv.excel
    return capped(table(csv.reader(io.StringIO(text), dialect)))


def word_text(path: Path) -> str:
    """A Word file's body: its paragraphs and tables, in the order they stand."""
    import docx
    from docx.table import Table

    document = docx.Document(str(path))
    blocks: list[str] = []
    for block in document.iter_inner_content():
        if isinstance(block, Table):
            text = table([[cell.text for cell in row.cells] for row in block.rows])
        else:
            text = " ".join(block.text.split())
        if text:
            blocks.append(text)
    return capped("\n\n".join(blocks))


def plain_text(path: Path) -> str:
    """A text or Markdown file as it is."""
    return capped(decoded(path.read_bytes()).strip())


def _clean(cell: str) -> str:
    return " ".join(cell.split()).replace("|", "\\|")


def _row(cells: Sequence[str], width: int) -> str:
    padded = [*cells[:width], *([""] * (width - len(cells)))]
    return "| " + " | ".join(padded) + " |"
