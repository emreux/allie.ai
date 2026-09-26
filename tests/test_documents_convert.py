"""`documents/convert.py` (plan.md D35): an Excel workbook, a CSV, a Word
file or a text file as the text the document model reads.

The files are written here, into the test's own folder, with the libraries
that read them; nothing reads the user's documents. Plain on purpose (the
owner, 2026-09-25, findings 8 and 10 of the review): a cell is its raw
value, and a Word file is its body without its header.
"""

from __future__ import annotations

import codecs
from datetime import datetime
from pathlib import Path
from typing import Any

import docx
import openpyxl
import pytest

from allie.documents import convert
from allie.documents.convert import (
    CUT,
    ROWS_CUT,
    capped,
    cell_text,
    csv_text,
    excel_text,
    plain_text,
    table,
    word_text,
)

# --------------------------------------------------------------------------
# Cells and tables
# --------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("value", "shown"),
    [
        (0.18, "0.18"),
        (12500.5, "12500.5"),
        (datetime(2026, 8, 14), "2026-08-14 00:00:00"),
        (True, "True"),
        (None, ""),
        ("  Ali   Yılmaz \n", "Ali Yılmaz"),
    ],
)
def test_a_cell_is_its_raw_value_on_one_line(value: object, shown: str) -> None:
    assert cell_text(value) == shown


def test_rows_become_a_table_under_their_first_row() -> None:
    rows = [["Hesap", "Bakiye", ""], ["", " ", ""], ["Kasa", "36250.75", ""]]

    assert table(rows) == "| Hesap | Bakiye |\n| --- | --- |\n| Kasa | 36250.75 |"


def test_a_short_row_is_padded_and_a_bar_does_not_end_a_cell() -> None:
    assert table([["a", "b"], ["x|y"]]) == "| a | b |\n| --- | --- |\n| x\\|y |  |"


def test_no_rows_is_no_table() -> None:
    assert table([["", " "], []]) == ""


def test_a_long_table_is_cut_and_says_so(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(convert, "MAX_ROWS", 2)

    text = table([["n"], ["1"], ["2"], ["3"]])

    assert text.splitlines() == [
        "| n |",
        "| --- |",
        "| 1 |",
        "| 2 |",
        ROWS_CUT.format(total=3, limit=2),
    ]


def test_a_long_file_is_cut_and_says_so(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(convert, "MAX_CHARS", 5)

    assert capped("abcdefgh") == "abcde\n" + CUT.format(total=8, limit=5)
    assert capped("abc") == "abc"


# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------


def test_a_workbook_is_one_table_a_sheet_and_an_empty_sheet_is_none(tmp_path: Path) -> None:
    book: Any = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Fatura"
    sheet.append(["Kalem", "Tutar", "KDV oranı", "Tarih", "Toplam"])
    sheet.append(["Danışmanlık", 12500.5, 0.18, datetime(2026, 8, 14), "=B2*(1+C2)"])
    book.create_sheet("Boş")
    path = tmp_path / "fatura.xlsx"
    book.save(path)

    assert excel_text(path) == (
        "## Fatura\n\n"
        "| Kalem | Tutar | KDV oranı | Tarih | Toplam |\n"
        "| --- | --- | --- | --- | --- |\n"
        # Raw values; a formula Excel never calculated has no value to show.
        "| Danışmanlık | 12500.5 | 0.18 | 2026-08-14 00:00:00 |  |"
    )


# --------------------------------------------------------------------------
# CSV
# --------------------------------------------------------------------------


def test_a_semicolon_file_in_the_turkish_code_page_is_read(tmp_path: Path) -> None:
    path = tmp_path / "kdv.csv"
    path.write_bytes("Dönem;Matrah;KDV\nAğustos;12.500,50;2.250,09\n".encode("cp1254"))

    assert csv_text(path) == (
        "| Dönem | Matrah | KDV |\n| --- | --- | --- |\n| Ağustos | 12.500,50 | 2.250,09 |"
    )


def test_a_comma_file_with_a_byte_order_mark_is_read(tmp_path: Path) -> None:
    path = tmp_path / "liste.csv"
    path.write_bytes(codecs.BOM_UTF8 + b"Ad,Tutar\nAli,100\n")

    assert csv_text(path) == "| Ad | Tutar |\n| --- | --- |\n| Ali | 100 |"


def test_a_single_column_is_still_a_table(tmp_path: Path) -> None:
    path = tmp_path / "tek.csv"
    path.write_text("Ad\nAli\nVeli\n", encoding="utf-8")

    assert csv_text(path) == "| Ad |\n| --- |\n| Ali |\n| Veli |"


# --------------------------------------------------------------------------
# Word and text
# --------------------------------------------------------------------------


def test_a_word_file_is_its_body_in_order_without_its_header(tmp_path: Path) -> None:
    document: Any = docx.Document()
    document.sections[0].header.paragraphs[0].text = "Bilgi Birikim A.Ş."
    document.add_paragraph("Sayın Ali Yılmaz,")
    grid = document.add_table(rows=2, cols=2)
    cells = {(0, 0): "Kalem", (0, 1): "Tutar", (1, 0): "Danışmanlık", (1, 1): "12.500,50 TL"}
    for (row, column), text in cells.items():
        grid.cell(row, column).text = text
    document.add_paragraph("Son ödeme: 30.09.2026")
    path = tmp_path / "yazi.docx"
    document.save(str(path))

    assert word_text(path) == (
        "Sayın Ali Yılmaz,\n\n"
        "| Kalem | Tutar |\n| --- | --- |\n| Danışmanlık | 12.500,50 TL |\n\n"
        "Son ödeme: 30.09.2026"
    )


def test_a_text_file_in_the_turkish_code_page_is_read(tmp_path: Path) -> None:
    path = tmp_path / "not.txt"
    path.write_bytes("Beyanname: 26 Eylül\n".encode("cp1254"))

    assert plain_text(path) == "Beyanname: 26 Eylül"
